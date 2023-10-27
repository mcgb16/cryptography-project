import random
import string
import secrets
import func_db
import shutil
from fpdf import FPDF
from tkinter import filedialog
from docx import Document
import openpyxl

def verify_cpf(user_cpf):
    cpf_not_valid = 'CPF não é válido. Cheque a digitação e lembre-se: apenas números.'
    if user_cpf == '11111111111':
        return cpf_not_valid
    elif user_cpf == '22222222222':
        return cpf_not_valid
    if user_cpf == '33333333333':
        return cpf_not_valid
    elif user_cpf == '44444444444':
        return cpf_not_valid
    if user_cpf == '55555555555':
        return cpf_not_valid
    elif user_cpf == '66666666666':
        return cpf_not_valid
    if user_cpf == '77777777777':
        return cpf_not_valid
    elif user_cpf == '88888888888':
        return cpf_not_valid
    if user_cpf == '99999999999':
        return cpf_not_valid
    elif user_cpf == '00000000000':
        return cpf_not_valid

    cpf_digit_list = []    
    v1 = 10
    v2 = 11
    i = 0
    j = 0
    cpf_digit1_validate = 0
    cpf_digit2_validate = 0

    if len(user_cpf) != 11:
        return cpf_not_valid

    try:
        int(user_cpf)
    except:
        return cpf_not_valid
    
    for d in user_cpf:
        cpf_digit_list.append(int(d))

    while i < 9:
        cpf_digit1_validate += cpf_digit_list[i]*v1
        v1 -= 1
        i += 1

    cpf_digit1_validate = (cpf_digit1_validate*10)%11

    if cpf_digit1_validate == 10:
        cpf_digit1_validate = 0
    
    if cpf_digit1_validate == cpf_digit_list[9]:
        while j < 10:
            cpf_digit2_validate += cpf_digit_list[j]*v2
            v2 -= 1
            j += 1

        cpf_digit2_validate = (cpf_digit2_validate*10)%11

        if cpf_digit2_validate == 10:
            cpf_digit2_validate = 0
        
        if cpf_digit2_validate == cpf_digit_list[10]:
            return 'cpf valid'
        else:
            return cpf_not_valid
    else:
        return cpf_not_valid

def verify_len_input_cryp(cpf, name, text_to_encrypt):
    null_message = 'Por favor preencha todos os campos!'
    len_message_text_to_encrypt = 'O limite máximo para criptografia de textos é de 255 caracteres.'
    no_dir_message = "Selecione um arquivo a ser criptografado, por favor."

    if cpf == '':
        return null_message
    elif name == '':
        return null_message
    elif text_to_encrypt == '':
        return null_message
    elif len(text_to_encrypt) > 255:
        return len_message_text_to_encrypt
    elif text_to_encrypt == 'no dir':
        return no_dir_message
    else:
        return 'valid'

def generate_encrypted_data(text_to_crypt):
    crypted_list = []
    crypted_data = ''
    
    for d in text_to_crypt:
        crypt_mask_len = random.randint(128,256)
        crypted_list.append(''.join(random.choices(string.ascii_letters + string.digits + string.punctuation, k = crypt_mask_len)))
    
    for i in crypted_list:
        crypted_data += i
    
    return crypted_data

def generate_unique_key():
    unique_key_lenght = 128
    key = ''
    for i in range(unique_key_lenght):
        key += ''.join(secrets.choice(string.ascii_letters + string.digits))
    return key

def generate_pdf_files(main_txt, pdf_type, user_name, file_name):
    pdf_file = FPDF()
    pdf_file.add_page()
    cell_x = 190
    page_h = pdf_file.h

    if pdf_type == 'key':
        header_text = f'{user_name}, muito obrigado por criptografar conosco! Segue abaixo sua Key única para descriptografia futura.'
        footer_text = 'Lembre-se: não a perca de forma alguma! Caso contrário, não será possível efetuar a descriptografia do texto gerado.'
        pdf_file.set_font("courier", size=12,style="I")
        pdf_file.multi_cell(cell_x,txt=header_text, align="C")
        pdf_file.set_y(page_h/2)
        pdf_file.set_font("courier", size=11)
        pdf_file.multi_cell(cell_x,txt=main_txt, align="L")  
    elif pdf_type == 'encrypted_text':
        footer_text = f'{user_name}, para efetuar a descriptografia, utilize a key fornecida, junto do CPF cadastrado, no aplicativo.'
        pdf_file.set_font("courier", size=11)
        pdf_file.multi_cell(cell_x,txt=main_txt, align="L")
    elif pdf_type == 'encryp_file':
        pdf_file.set_font("courier", size=11)
        pdf_file.multi_cell(cell_x,txt=main_txt, align="L")
        pdf_file.output(file_name)
        
        return "Sucesso!"

    pdf_file.set_y(page_h - pdf_file.t_margin - pdf_file.b_margin)
    pdf_file.set_font("courier", size=12,style="I")
    pdf_file.multi_cell(cell_x,txt=footer_text, align="C")

    pdf_file.output(filedialog.asksaveasfilename(defaultextension=".pdf", filetypes=[("All Files", "*.*")],initialfile=file_name))

    return 'Sucesso!'

def pdf_files_controller(encrypted_text, key, user_name, cryp_type, file_name=''):
    if cryp_type == 'text':
        generate_pdf_files(encrypted_text, 'encrypted_text', user_name, 'Encrypted_text')
        generate_pdf_files(key, 'key', user_name, 'Key')
    elif cryp_type == 'file':
        generate_pdf_files(key, 'key', user_name, 'Key')
    elif cryp_type == 'encryp_file':
        generate_pdf_files(encrypted_text,'encryp_file', user_name, file_name)

    return 'Todos os pdfs criados com sucesso.' 

def send_text_to_db(key, name, text, cpf):
    send_to_db = func_db.save_text_on_db(key,name,text,cpf)
    
    return send_to_db

def search_file():
    file_dir = filedialog.askopenfilename(initialdir='/', filetypes=[("All Files", "*.*")])
    file_dir_list = file_dir.split('/')
    file_name = file_dir_list[-1]

    return file_dir, file_name

def generate_encrypted_file(encrypted_data, file_name, file_dir):
    success_msg = 'Arquivo criptografado com êxito!'
    
    if '.docx' in file_name:
        docx_file = Document()
        docx_file.add_paragraph(encrypted_data)
        docx_file.save(file_dir)
        return success_msg
    elif '.xlsx' in file_name:
        wb = openpyxl.Workbook()
        encrypted_sheet = wb.active
        
        encrypted_data_len = len(encrypted_data)

        for i in range(1,200):

            random_len = random.randint(encrypted_data_len//2, encrypted_data_len)
            encrypted_sheet[f'A{i}'] = encrypted_data[random_len:]

        wb.save(file_dir)

        return success_msg
    elif '.pdf' in file_name:
        pdf_file = pdf_files_controller(encrypted_data, '','','encryp_file', file_dir)
        return success_msg
    elif '.txt' or '.md' in file_name:
        with open(file_dir,'w') as file:
            file.write(encrypted_data)
        return success_msg

def move_file_to_server(file_dir, file_name):   
    server_dir = 'C:\\Users\\mathe\\Desktop\\sv'
    try:
        shutil.move(file_dir,server_dir)
        server_file_dir = server_dir + "\\" + file_name
        server_file_dir = server_file_dir.replace("\\", "/")
        return server_file_dir
    except:
        print('Erro ao mover o arquivo para o servidor.')

def send_file_to_db(key, name, cpf, file_name, file_dir):
    send_to_db = func_db.save_file_on_db(key, name, cpf, file_name, file_dir)
    
    return send_to_db

def search_on_db(key, cpf, decryp_type):
    if decryp_type == 'text':
        verify_db = func_db.search_text_on_db(key, cpf)
        return verify_db
    elif decryp_type == 'file':
        verify_db = func_db.search_file_on_db(key, cpf)
        return verify_db

def verify_len_input_decryp(cpf, key, picklist):
    null_message = 'Por favor preencha todos os campos!'
    null_picklist_message = 'Selecione um tipo de descriptografia.'
    
    if cpf == '':
        return null_message
    elif key == '':
        return null_message
    elif picklist == '':
        return null_picklist_message
    else:
        return 'valid'

def save_decrypted_file(file_name, file_dir):
    final_dir_name = filedialog.asksaveasfilename(filetypes=[("All Files", "*.*")],initialfile=file_name)
    final_dir_list = final_dir_name.split("/")

    final_dir = ''

    for i in final_dir_list:
        if i != final_dir_list[-1]:
            final_dir += i + "\\"

    shutil.move(file_dir, final_dir)
    return